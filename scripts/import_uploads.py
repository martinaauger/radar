#!/usr/bin/env python3
"""Import Laura (LF) Excel + LinkedIn Connections CSV into Radar SQLite."""

from __future__ import annotations

import csv
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("/workspace")
DB_PATH = ROOT / "data" / "radar.sqlite"
UPLOADS = Path("/home/ubuntu/.cursor/projects/workspace/uploads")
LF_XLSX = UPLOADS / "Perfil_LF-_Base_de_contactos_567e.xlsx"
CONNECTIONS = UPLOADS / "Connections_bcd5.csv"

NOW = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")


def uid() -> str:
    return str(uuid.uuid4())


def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() in {"sin dato", "none", "nan", "-"}:
        return ""
    return s


def map_sector(raw: str) -> str:
    s = (raw or "").lower()
    if "agua" in s or "saneamiento" in s:
        # multi with agua still prefer primary listed first
        pass
    parts = [p.strip() for p in (raw or "").split("·")]
    primary = parts[0].lower() if parts else ""
    if "nuevo" in primary or "residuo" in primary:
        return "nuevos_negocios"
    if "generaci" in primary:
        return "generacion"
    if "gas" in primary:
        return "gas"
    if "agua" in primary or "saneamiento" in primary:
        return "agua"
    if "electric" in primary or "distribuci" in primary or "transmisi" in primary:
        return "electricidad"
    if "equipo" in primary or "tecnolog" in primary:
        return "transversal"
    if len(parts) > 1:
        # multi-sector → transversal
        return "transversal"
    return "transversal"


def map_org(tipo: str, sector: str) -> str:
    t = (tipo or "").lower()
    s = (sector or "").lower()
    if "regul" in t:
        return "regulador"
    if "asocia" in t or "cámara" in t or "camara" in t:
        return "asociacion"
    if "público" in t or "publico" in t or "gobierno" in t:
        return "gobierno"
    if "independ" in t:
        return "otro"
    if "empresa" in t:
        if "generaci" in s:
            return "generadora"
        if "gas" in s:
            return "distribuidora"
        if "agua" in s:
            return "distribuidora"
        if "transmisi" in s or "transport" in s:
            return "transportista"
        if "electric" in s or "distribuci" in s:
            return "distribuidora"
        return "otro"
    return "otro"


def map_decision(puesto: str, cargo: str) -> str:
    text = f"{puesto} {cargo}".lower()
    if any(x in text for x in ("ceo", "director", "gerente", "presidente", "cfo", "coo")):
        return "decisor"
    if any(x in text for x in ("asesor", "regul", "tarifa", "legal", "regulator")):
        return "influenciador"
    if any(x in text for x in ("especialista", "analista", "ingenier", "técnic", "tecnic")):
        return "tecnico"
    if any(x in text for x in ("coordina", "jefe", "head")):
        return "influenciador"
    return "desconocido"


def normalize_li(url: str) -> str:
    u = clean(url).lower().rstrip("/")
    u = u.split("?")[0]
    return u


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS companies (
          id TEXT PRIMARY KEY,
          name TEXT NOT NULL,
          country TEXT NOT NULL DEFAULT '',
          city TEXT,
          business_unit TEXT NOT NULL,
          org_type TEXT NOT NULL,
          notes TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS contacts (
          id TEXT PRIMARY KEY,
          company_id TEXT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
          full_name TEXT NOT NULL,
          title TEXT NOT NULL DEFAULT '',
          email TEXT,
          phone TEXT,
          linkedin_url TEXT,
          decision_level TEXT NOT NULL DEFAULT 'desconocido',
          mail_status TEXT NOT NULL DEFAULT 'pendiente',
          notes TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS contact_owners (
          contact_id TEXT NOT NULL REFERENCES contacts(id) ON DELETE CASCADE,
          owner TEXT NOT NULL,
          PRIMARY KEY (contact_id, owner)
        );
        CREATE TABLE IF NOT EXISTS contact_uses (
          contact_id TEXT NOT NULL REFERENCES contacts(id) ON DELETE CASCADE,
          use_tag TEXT NOT NULL,
          PRIMARY KEY (contact_id, use_tag)
        );
        CREATE TABLE IF NOT EXISTS contact_sources (
          id TEXT PRIMARY KEY,
          contact_id TEXT NOT NULL REFERENCES contacts(id) ON DELETE CASCADE,
          source_type TEXT NOT NULL,
          owner TEXT,
          note TEXT,
          recorded_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_companies_bu ON companies(business_unit);
        CREATE INDEX IF NOT EXISTS idx_contacts_company ON contacts(company_id);
        CREATE INDEX IF NOT EXISTS idx_contacts_li ON contacts(linkedin_url);
        """
    )


def reset_db(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        DELETE FROM contact_sources;
        DELETE FROM contact_uses;
        DELETE FROM contact_owners;
        DELETE FROM contacts;
        DELETE FROM companies;
        """
    )


class Importer:
    def __init__(self, conn: sqlite3.Connection):
        self.conn = conn
        self.company_cache: dict[tuple[str, str], str] = {}
        self.li_index: dict[str, str] = {}
        self.stats = {
            "companies": 0,
            "contacts_laura": 0,
            "contacts_connections": 0,
            "merged": 0,
            "with_email": 0,
        }

    def get_or_create_company(
        self,
        name: str,
        country: str,
        city: str,
        business_unit: str,
        org_type: str,
    ) -> str:
        key_name = name or "(Sin empresa)"
        key_country = country or ""
        cache_key = (key_name.lower(), key_country.lower())
        if cache_key in self.company_cache:
            return self.company_cache[cache_key]
        existing = self.conn.execute(
            "SELECT id FROM companies WHERE lower(name)=lower(?) AND lower(country)=lower(?)",
            (key_name, key_country),
        ).fetchone()
        if existing:
            self.company_cache[cache_key] = existing[0]
            return existing[0]
        cid = uid()
        self.conn.execute(
            """INSERT INTO companies
               (id, name, country, city, business_unit, org_type, notes, created_at, updated_at)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (
                cid,
                key_name,
                key_country,
                city or None,
                business_unit,
                org_type,
                None,
                NOW,
                NOW,
            ),
        )
        self.company_cache[cache_key] = cid
        self.stats["companies"] += 1
        return cid

    def add_owner(self, contact_id: str, owner: str) -> None:
        self.conn.execute(
            "INSERT OR IGNORE INTO contact_owners (contact_id, owner) VALUES (?,?)",
            (contact_id, owner),
        )

    def add_use(self, contact_id: str, use: str) -> None:
        self.conn.execute(
            "INSERT OR IGNORE INTO contact_uses (contact_id, use_tag) VALUES (?,?)",
            (contact_id, use),
        )

    def add_source(
        self, contact_id: str, source_type: str, owner: str | None, note: str
    ) -> None:
        self.conn.execute(
            """INSERT INTO contact_sources
               (id, contact_id, source_type, owner, note, recorded_at)
               VALUES (?,?,?,?,?,?)""",
            (uid(), contact_id, source_type, owner, note, NOW),
        )

    def insert_contact(
        self,
        *,
        company_id: str,
        full_name: str,
        title: str,
        email: str | None,
        linkedin_url: str | None,
        decision_level: str,
        owner: str,
        source_type: str,
        note: str,
        uses: list[str],
    ) -> str:
        cid = uid()
        mail_status = "completo" if email else "pendiente"
        self.conn.execute(
            """INSERT INTO contacts
               (id, company_id, full_name, title, email, phone, linkedin_url,
                decision_level, mail_status, notes, created_at, updated_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                cid,
                company_id,
                full_name,
                title,
                email,
                None,
                linkedin_url,
                decision_level,
                mail_status,
                None,
                NOW,
                NOW,
            ),
        )
        self.add_owner(cid, owner)
        for u in uses:
            self.add_use(cid, u)
        self.add_source(cid, source_type, owner, note)
        if linkedin_url:
            self.li_index[normalize_li(linkedin_url)] = cid
        if email:
            self.stats["with_email"] += 1
        return cid

    def import_laura(self) -> None:
        wb = load_workbook(LF_XLSX, read_only=True, data_only=True)
        ws = wb["Contactos"]
        first = True
        for row in ws.iter_rows(values_only=True):
            if first:
                first = False
                continue
            nombre, apellido = clean(row[0]), clean(row[1])
            if not nombre and not apellido:
                continue
            full_name = f"{nombre} {apellido}".strip()
            empresa = clean(row[2]) or "(Sin empresa)"
            cargo = clean(row[3])
            puesto = clean(row[4])
            pais = clean(row[5])
            ciudad = clean(row[6])
            sector_raw = clean(row[7])
            tipo = clean(row[8])
            email = clean(row[10]) or None
            li = clean(row[12]) or None

            bu = map_sector(sector_raw)
            org = map_org(tipo, sector_raw)
            decision = map_decision(puesto, cargo)
            company_id = self.get_or_create_company(
                empresa, pais, ciudad, bu, org
            )

            uses: list[str] = []
            if org == "regulador" or "regul" in f"{puesto} {cargo}".lower():
                uses.append("licitacion")
            if "socio" in sector_raw.lower() or tipo.lower() == "asociación":
                uses.append("asociado")

            self.insert_contact(
                company_id=company_id,
                full_name=full_name,
                title=cargo or puesto,
                email=email,
                linkedin_url=li,
                decision_level=decision,
                owner="laura",
                source_type="manual",
                note="Base Laura (Perfil LF)",
                uses=uses,
            )
            self.stats["contacts_laura"] += 1
        wb.close()

    def import_connections(self) -> None:
        raw = CONNECTIONS.read_text(encoding="utf-8", errors="replace")
        lines = raw.splitlines()
        header_idx = next(
            i for i, line in enumerate(lines) if line.startswith("First Name")
        )
        reader = csv.DictReader(lines[header_idx:])
        for row in reader:
            first = clean(row.get("First Name"))
            last = clean(row.get("Last Name"))
            if not first and not last:
                continue
            full_name = f"{first} {last}".strip()
            company = clean(row.get("Company")) or "(Sin empresa)"
            title = clean(row.get("Position"))
            email = clean(row.get("Email Address")) or None
            li = clean(row.get("URL")) or None
            li_key = normalize_li(li) if li else ""

            # Infer sector lightly from company/title keywords
            blob = f"{company} {title}".lower()
            if any(x in blob for x in ("gas", "gnl", "metrogas", "enargas")):
                bu = "gas"
            elif any(
                x in blob
                for x in ("água", "agua", "saneamiento", "saneamento", "aysa")
            ):
                bu = "agua"
            elif any(
                x in blob
                for x in (
                    "gera",
                    "genera",
                    "abrage",
                    "abiape",
                    "statkraft",
                    "aes",
                )
            ):
                bu = "generacion"
            elif any(
                x in blob
                for x in (
                    "energi",
                    "electric",
                    "distrib",
                    "aneel",
                    "enre",
                    "copel",
                    "cemig",
                    "transener",
                    "isa ",
                )
            ):
                bu = "electricidad"
            else:
                bu = "transversal"

            org = "asociacion" if any(
                x in blob for x in ("associa", "asocia", "câmara", "camara", "abrage", "abiape")
            ) else ("regulador" if any(x in blob for x in ("aneel", "enre", "enargas", "osinerg", "creg", "regul")) else "otro")

            if li_key and li_key in self.li_index:
                cid = self.li_index[li_key]
                self.add_owner(cid, "fernando")
                self.add_source(
                    cid,
                    "linkedin_export",
                    "fernando",
                    "Connections.csv — merge por LinkedIn",
                )
                if email:
                    self.conn.execute(
                        """UPDATE contacts SET email=COALESCE(email, ?),
                           mail_status=CASE WHEN email IS NULL OR email='' THEN
                             CASE WHEN ? IS NOT NULL THEN 'completo' ELSE mail_status END
                           ELSE mail_status END,
                           updated_at=? WHERE id=?""",
                        (email, email, NOW, cid),
                    )
                self.stats["merged"] += 1
                continue

            company_id = self.get_or_create_company(
                company, "", "", bu, org
            )
            self.insert_contact(
                company_id=company_id,
                full_name=full_name,
                title=title,
                email=email,
                linkedin_url=li,
                decision_level=map_decision("", title),
                owner="fernando",
                source_type="linkedin_export",
                note="Connections.csv",
                uses=[],
            )
            self.stats["contacts_connections"] += 1


def main() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    ensure_schema(conn)
    reset_db(conn)
    imp = Importer(conn)
    print("Importing Laura (LF)…")
    imp.import_laura()
    print("Importing Connections.csv…")
    imp.import_connections()
    conn.commit()
    companies = conn.execute("SELECT COUNT(*) FROM companies").fetchone()[0]
    contacts = conn.execute("SELECT COUNT(*) FROM contacts").fetchone()[0]
    emails = conn.execute(
        "SELECT COUNT(*) FROM contacts WHERE mail_status='completo'"
    ).fetchone()[0]
    regul = conn.execute(
        """SELECT COUNT(*) FROM contacts ct
           JOIN companies c ON c.id = ct.company_id
           WHERE c.org_type='regulador'
              OR lower(ct.title) LIKE '%regul%'
              OR lower(ct.title) LIKE '%tarifa%'"""
    ).fetchone()[0]
    print("DONE")
    print("stats", imp.stats)
    print(
        f"db companies={companies} contacts={contacts} emails={emails} regulation≈{regul}"
    )
    conn.close()


if __name__ == "__main__":
    main()
